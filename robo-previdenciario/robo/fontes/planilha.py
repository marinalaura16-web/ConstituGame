"""Leitura genérica de CSV/XLSX com mapeamento de colunas por apelidos."""
from __future__ import annotations

import csv
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook

from ..texto import sem_acento


def _chave(s: object) -> str:
    return " ".join(sem_acento(str(s or "")).lower().replace("_", " ").split())


def ler_linhas(caminho: Path, aba: str | None = None) -> list[dict]:
    """Devolve as linhas como dicionários {cabeçalho: valor}."""
    caminho = Path(caminho)
    if caminho.suffix.lower() in (".csv", ".txt"):
        bruto = caminho.read_bytes()
        try:
            texto = bruto.decode("utf-8-sig")
        except UnicodeDecodeError:
            texto = bruto.decode("latin-1")  # exportações de sistemas brasileiros
        dialeto = csv.Sniffer().sniff(texto[:5000], delimiters=";,\t")
        return list(csv.DictReader(texto.splitlines(), dialect=dialeto))
    wb = load_workbook(caminho, read_only=True, data_only=True)
    ws = wb[aba] if aba else wb.worksheets[0]
    linhas = ws.iter_rows(values_only=True)
    cabecalho = [str(c or "").strip() for c in next(linhas, [])]
    saida = [dict(zip(cabecalho, valores)) for valores in linhas if any(v not in (None, "") for v in valores)]
    wb.close()
    return saida


def mapear(linha: dict, apelidos: dict[str, list[str]]) -> dict:
    """Troca os cabeçalhos da planilha pelos nomes internos definidos em planilhas.yaml."""
    por_chave = {_chave(k): v for k, v in linha.items()}
    saida = {}
    for campo, nomes in apelidos.items():
        valor = None
        for nome in nomes:
            if _chave(nome) in por_chave:
                valor = por_chave[_chave(nome)]
                break
        if isinstance(valor, str):
            valor = valor.strip() or None
        saida[campo] = valor
    return saida


def para_data(valor: object) -> date | None:
    if valor in (None, ""):
        return None
    if isinstance(valor, datetime):
        return valor.date()
    if isinstance(valor, date):
        return valor
    s = str(valor).strip()[:10]
    for formato in ("%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y", "%d/%m/%y"):
        try:
            return datetime.strptime(s, formato).date()
        except ValueError:
            continue
    raise ValueError(f"data inválida: {valor!r}")
